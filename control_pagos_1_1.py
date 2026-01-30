"""
AUTOMATIZACIÓN COMPLETA - CONTROL DE PAGOS - VERSIÓN CORREGIDA
Corrige problemas con macros que ocultan hojas y manejo de columnas
"""

import pandas as pd
import win32com.client
import pythoncom
from openpyxl.styles import Font, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
from pathlib import Path
from datetime import datetime, timedelta
import locale
import tkinter as tk
from tkinter import ttk, messagebox
from tkcalendar import DateEntry
import configparser
import sys
import time

# Configuración de español
try:
    locale.setlocale(locale.LC_TIME, 'Spanish_Spain.1252')  # Windows
except locale.Error:
    try:
        locale.setlocale(locale.LC_TIME, 'es_ES.UTF-8')  # Linux
    except locale.Error:
        pass

class InterfazModerna:
    """
    Interfaz gráfica moderna para seleccionar la fecha de filtrado (Proyección)
    """
    def __init__(self):
        self.fecha_seleccionada = None
        self.ejecutar_proceso = False
        
        # Colores del tema
        self.COLOR_PRIMARIO = "#2C3E50"
        self.COLOR_SECUNDARIO = "#3498DB"
        self.COLOR_ACENTO = "#27AE60"
        self.COLOR_FONDO = "#ECF0F1"
        self.COLOR_TEXTO = "#2C3E50"
        self.COLOR_ERROR = "#E74C3C"
        
    def crear_ventana(self):
        """Crea la ventana de interfaz moderna"""
        self.root = tk.Tk()
        self.root.title("Control de Pagos GCO")
        self.root.geometry("700x700")
        self.root.resizable(False, False)
        self.root.configure(bg=self.COLOR_FONDO)
        
        # Centrar ventana
        self.centrar_ventana()
        
        # Configurar estilo
        self.configurar_estilos()
        
        # Frame principal con gradiente simulado
        main_frame = tk.Frame(self.root, bg=self.COLOR_FONDO)
        main_frame.pack(fill=tk.BOTH, expand=True, padx=0, pady=0)
        
        # Header con color
        self.crear_header(main_frame)
        
        # Contenido principal
        self.crear_contenido(main_frame)
        
        # Footer con botones
        self.crear_footer(main_frame)
        
        # Agregar icono si existe
        try:
            self.root.iconbitmap('icon.ico')
        except:
            pass
        
        self.root.mainloop()
    
    def configurar_estilos(self):
        """Configura los estilos personalizados"""
        style = ttk.Style()
        style.theme_use('clam')
        
        # Estilo para LabelFrame
        style.configure(
            "Modern.TLabelframe",
            background=self.COLOR_FONDO,
            bordercolor=self.COLOR_SECUNDARIO,
            borderwidth=2
        )
        style.configure(
            "Modern.TLabelframe.Label",
            background=self.COLOR_FONDO,
            foreground=self.COLOR_PRIMARIO,
            font=("Segoe UI", 11, "bold")
        )
        
        # Estilo para Labels
        style.configure(
            "Title.TLabel",
            background=self.COLOR_PRIMARIO,
            foreground="white",
            font=("Segoe UI", 20, "bold")
        )
        
        style.configure(
            "Subtitle.TLabel",
            background=self.COLOR_PRIMARIO,
            foreground="white",
            font=("Segoe UI", 11)
        )
    
    def crear_header(self, parent):
        """Crea el header con título y logo"""
        header_frame = tk.Frame(parent, bg=self.COLOR_PRIMARIO, height=140)
        header_frame.pack(fill=tk.X, pady=0)
        header_frame.pack_propagate(False)
        
        # Contenedor centrado
        content = tk.Frame(header_frame, bg=self.COLOR_PRIMARIO)
        content.place(relx=0.5, rely=0.5, anchor="center")
        
        # Icono (emoji como placeholder)
        icon_label = tk.Label(
            content,
            text="📊",
            font=("Segoe UI", 30),
            bg=self.COLOR_PRIMARIO,
            fg="white"
        )
        icon_label.pack(side=tk.LEFT, padx=(0, 15))
        
        # Textos
        text_frame = tk.Frame(content, bg=self.COLOR_PRIMARIO)
        text_frame.pack(side=tk.LEFT)
        
        titulo = tk.Label(
            text_frame,
            text="Control de Pagos",
            font=("Segoe UI", 20, "bold"),
            bg=self.COLOR_PRIMARIO,
            fg="white"
        )
        titulo.pack(anchor="w")
        
        subtitulo = tk.Label(
            text_frame,
            text="Sistema de Gestión de Importaciones",
            font=("Segoe UI", 11),
            bg=self.COLOR_PRIMARIO,
            fg="#BDC3C7"
        )
        subtitulo.pack(anchor="w")
    
    def crear_contenido(self, parent):
        """Crea el contenido principal"""
        content_frame = tk.Frame(parent, bg=self.COLOR_FONDO)
        content_frame.pack(fill=tk.BOTH, expand=True, padx=30, pady=30)
        
        # Tarjeta principal
        card_frame = tk.Frame(
            content_frame,
            bg="white",
            relief=tk.FLAT,
            borderwidth=0
        )
        card_frame.pack(fill=tk.BOTH, expand=True)
        
        # Agregar sombra simulada con bordes
        self.agregar_sombra(card_frame)
        
        # Padding interno
        inner_frame = tk.Frame(card_frame, bg="white")
        inner_frame.pack(fill=tk.BOTH, expand=True, padx=30, pady=30)
        
        # Título de la sección
        section_title = tk.Label(
            inner_frame,
            text="📅 Selección de Fecha de Proyección",
            font=("Segoe UI", 14, "bold"),
            bg="white",
            fg=self.COLOR_PRIMARIO
        )
        section_title.pack(pady=(0, 5))
        
        # Línea separadora
        separator = tk.Frame(inner_frame, height=2, bg=self.COLOR_SECUNDARIO)
        separator.pack(fill=tk.X, pady=(0, 20))
        
        # Descripción
        desc_label = tk.Label(
            inner_frame,
            text="Selecciona la fecha para la cual deseas generar la proyección de pagos.\nPor defecto, se sugiere el próximo miércoles.",
            font=("Segoe UI", 10),
            bg="white",
            fg="#7F8C8D",
            justify=tk.CENTER
        )
        desc_label.pack(pady=(0, 25))
        
        # Frame para el calendario
        cal_frame = tk.Frame(inner_frame, bg="white")
        cal_frame.pack(pady=10)
        
        # DateEntry con estilo mejorado
        self.calendario = DateEntry(
            cal_frame,
            width=22,
            background=self.COLOR_SECUNDARIO,
            foreground='white',
            borderwidth=2,
            font=("Segoe UI", 12),
            date_pattern='dd/mm/yyyy',
            locale='es_ES',
            selectbackground=self.COLOR_ACENTO,
            selectforeground='white'
        )
        self.calendario.pack(pady=10)
        
        # Calcular próximo miércoles por defecto
        proximo_miercoles = self.obtener_proximo_miercoles(datetime.now())
        self.calendario.set_date(proximo_miercoles)
        
        # Frame para información de fecha
        info_frame = tk.Frame(inner_frame, bg="white")
        info_frame.pack(pady=15)
        
        # Mostrar día de la semana seleccionado
        self.dia_semana_label = tk.Label(
            info_frame,
            text="",
            font=("Segoe UI", 12, "bold"),
            bg="white"
        )
        self.dia_semana_label.pack()
        
        # Actualizar día de la semana
        self.actualizar_dia_semana()
        self.calendario.bind("<<DateEntrySelected>>", lambda e: self.actualizar_dia_semana())
        
        # Nota informativa
        note_frame = tk.Frame(inner_frame, bg="#E8F8F5", relief=tk.FLAT, borderwidth=1)
        note_frame.pack(fill=tk.X, pady=(20, 0))
        
        note_icon = tk.Label(
            note_frame,
            text="ℹ️",
            font=("Segoe UI", 14),
            bg="#E8F8F5"
        )
        note_icon.pack(side=tk.LEFT, padx=10, pady=10)
        
        note_text = tk.Label(
            note_frame,
            text="Se recomienda seleccionar miércoles para las proyecciones semanales",
            font=("Segoe UI", 9),
            bg="#E8F8F5",
            fg="#16A085",
            justify=tk.LEFT
        )
        note_text.pack(side=tk.LEFT, pady=10, padx=(0, 10))
    
    def crear_footer(self, parent):
        """Crea el footer con botones de acción"""
        footer_frame = tk.Frame(parent, bg=self.COLOR_FONDO, height=70)
        footer_frame.pack(fill=tk.X, side=tk.BOTTOM, padx=30, pady=(0, 20))
        footer_frame.pack_propagate(False)
        
        # Contenedor de botones
        button_container = tk.Frame(footer_frame, bg=self.COLOR_FONDO)
        button_container.place(relx=0.5, rely=0.5, anchor="center")
        
        # Botón Ejecutar
        self.btn_ejecutar = tk.Button(
            button_container,
            text="▶  EJECUTAR PROCESO",
            command=self.ejecutar,
            bg=self.COLOR_ACENTO,
            fg="white",
            font=("Segoe UI", 11, "bold"),
            width=20,
            height=2,
            cursor="hand2",
            relief=tk.FLAT,
            borderwidth=0
        )
        self.btn_ejecutar.pack(side=tk.LEFT, padx=10)
        
        # Botón Cancelar
        self.btn_cancelar = tk.Button(
            button_container,
            text="✕  CANCELAR",
            command=self.cancelar,
            bg=self.COLOR_ERROR,
            fg="white",
            font=("Segoe UI", 11),
            width=15,
            height=2,
            cursor="hand2",
            relief=tk.FLAT,
            borderwidth=0
        )
        self.btn_cancelar.pack(side=tk.LEFT, padx=10)
        
        # Efectos hover con animación suave
        self.agregar_efectos_hover(self.btn_ejecutar, self.COLOR_ACENTO, "#229954")
        self.agregar_efectos_hover(self.btn_cancelar, self.COLOR_ERROR, "#C0392B")
    
    def agregar_sombra(self, widget):
        """Simula sombra en un widget"""
        shadow = tk.Frame(
            widget.master,
            bg="#95A5A6",
            relief=tk.FLAT
        )
        shadow.place(in_=widget, x=3, y=3, relwidth=1, relheight=1)
        widget.lift()
    
    def agregar_efectos_hover(self, boton, color_normal, color_hover):
        """Agrega efectos hover a los botones"""
        def on_enter(e):
            boton.config(bg=color_hover)
            
        def on_leave(e):
            boton.config(bg=color_normal)
        
        boton.bind("<Enter>", on_enter)
        boton.bind("<Leave>", on_leave)
    
    def centrar_ventana(self):
        """Centra la ventana en la pantalla"""
        self.root.update_idletasks()
        width = self.root.winfo_width()
        height = self.root.winfo_height()
        x = (self.root.winfo_screenwidth() // 2) - (width // 2)
        y = (self.root.winfo_screenheight() // 2) - (height // 2)
        self.root.geometry(f'{width}x{height}+{x}+{y}')
    
    def actualizar_dia_semana(self):
        """Actualiza el label con el día de la semana seleccionado"""
        fecha = self.calendario.get_date()
        dias_semana = ['Lunes', 'Martes', 'Miércoles', 'Jueves', 'Viernes', 'Sábado', 'Domingo']
        dia = dias_semana[fecha.weekday()]
        
        if fecha.weekday() == 2:  # Miércoles
            self.dia_semana_label.config(
                text=f"✓ {dia} {fecha.strftime('%d/%m/%Y')}",
                foreground=self.COLOR_ACENTO
            )
        else:
            self.dia_semana_label.config(
                text=f"{dia} {fecha.strftime('%d/%m/%Y')}",
                foreground=self.COLOR_SECUNDARIO
            )
    
    def obtener_proximo_miercoles(self, fecha):
        """Calcula el próximo miércoles"""
        dias_hasta_miercoles = (2 - fecha.weekday()) % 7
        if dias_hasta_miercoles == 0:
            dias_hasta_miercoles = 7
        return fecha + timedelta(days=dias_hasta_miercoles)
    
    def ejecutar(self):
        """Ejecuta el proceso"""
        self.fecha_seleccionada = self.calendario.get_date()
        self.ejecutar_proceso = True
        self.root.destroy()
    
    def cancelar(self):
        """Cancela el proceso"""
        if messagebox.askyesno("Confirmar", "¿Estás seguro de que deseas cancelar?"):
            self.ejecutar_proceso = False
            self.root.destroy()

class VentanaProgreso:
    """Ventana moderna de progreso"""
    def __init__(self, parent=None):
        self.ventana = tk.Toplevel(parent) if parent else tk.Tk()
        self.ventana.title("Procesando...")
        self.ventana.geometry("500x350")
        self.ventana.resizable(False, False)
        self.ventana.configure(bg="#ECF0F1")
        
        # Centrar
        self.centrar_ventana()
        
        # Frame principal
        main_frame = tk.Frame(self.ventana, bg="white")
        main_frame.pack(fill=tk.BOTH, expand=True, padx=20, pady=20)
        
        # Título
        titulo = tk.Label(
            main_frame,
            text="⚙️ Procesando Control de Pagos",
            font=("Segoe UI", 14, "bold"),
            bg="white",
            fg="#2C3E50"
        )
        titulo.pack(pady=(20, 10))
        
        # Mensaje
        self.mensaje_label = tk.Label(
            main_frame,
            text="Iniciando proceso...",
            font=("Segoe UI", 10),
            bg="white",
            fg="#7F8C8D"
        )
        self.mensaje_label.pack(pady=10)
        
        # Barra de progreso
        self.progreso = ttk.Progressbar(
            main_frame,
            length=400,
            mode='indeterminate'
        )
        self.progreso.pack(pady=20)
        self.progreso.start(10)
        
        # Log de acciones
        self.log_text = tk.Text(
            main_frame,
            height=10,
            width=50,
            font=("Consolas", 8),
            bg="#F8F9F9",
            fg="#2C3E50",
            relief=tk.FLAT
        )
        self.log_text.pack(pady=(0, 20), padx=20)
        self.log_text.config(state=tk.DISABLED)
    
    def actualizar_mensaje(self, mensaje):
        """Actualiza el mensaje de progreso"""
        self.mensaje_label.config(text=mensaje)
        self.ventana.update()
    
    def agregar_log(self, mensaje):
        """Agrega una línea al log"""
        self.log_text.config(state=tk.NORMAL)
        self.log_text.insert(tk.END, f"• {mensaje}\n")
        self.log_text.see(tk.END)
        self.log_text.config(state=tk.DISABLED)
        self.ventana.update()
    
    def centrar_ventana(self):
        """Centra la ventana"""
        self.ventana.update_idletasks()
        width = self.ventana.winfo_width()
        height = self.ventana.winfo_height()
        x = (self.ventana.winfo_screenwidth() // 2) - (width // 2)
        y = (self.ventana.winfo_screenheight() // 2) - (height // 2)
        self.ventana.geometry(f'{width}x{height}+{x}+{y}')
    
    def cerrar(self):
        """Cierra la ventana"""
        self.progreso.stop()
        self.ventana.destroy()

class CopiarArchivo:
    """Clase principal para el procesamiento de archivos - VERSIÓN CORREGIDA"""
    def __init__(self, fecha_filtrado=None, ventana_progreso=None):
        # Configuración de rutas
        self.config = configparser.ConfigParser()
        
        # Determinar ubicación del ejecutable o script
        if getattr(sys, 'frozen', False):
            application_path = Path(sys.executable).parent
        else:
            application_path = Path(__file__).parent
            
        config_path = application_path / 'config.ini'
        
        rutas_configuradas = False
        
        if config_path.exists():
            try:
                self.config.read(config_path, encoding='utf-8')
                if 'RUTAS' in self.config:
                    self.ruta_origen = Path(self.config['RUTAS'].get('ArchivoOrigen', ''))
                    self.ruta_intermedio = Path(self.config['RUTAS'].get('CarpetaIntermedia', ''))
                    self.ruta_destino_final = Path(self.config['RUTAS'].get('ArchivoFinal', ''))
                    rutas_configuradas = True
            except Exception as e:
                print(f"Error leyendo config.ini: {e}")

        # Si no hay config, usar rutas por defecto
        if not rutas_configuradas:
            base_path = Path.home() / "OneDrive - GCO" / "Escritorio"
            
            if not base_path.exists():
                 base_path_alt = Path.home() / "OneDrive" / "Escritorio"
                 if base_path_alt.exists():
                     base_path = base_path_alt
                 else:
                     base_path = Path.home() / "Desktop"
                     if not base_path.exists():
                         base_path = Path.home() / "Escritorio"
            
            self.ruta_origen = base_path / "00.CONTROL DE PAGOS 2026 1.xlsm"
            self.ruta_intermedio = base_path / "finanzas" / "info bancos" / "Pagos internacionales" / "proyección semana"
            self.ruta_destino_final = base_path / "finanzas" / "info bancos" / "Pagos internacionales" / "CONTROL PAGOS.xlsx"

        # NOMBRES DE HOJAS
        self.nombre_primera_hoja = "Control_Pagos"
        
        # FECHA DE PROYECCIÓN
        self.fecha_filtrado = fecha_filtrado
        
        # Ventana de progreso
        self.ventana_progreso = ventana_progreso
        
        # COLUMNAS PARA LA SEGUNDA HOJA
        self.columnas_segunda_hoja = [
            'IMPORTADOR',
            'MARCA', 
            'PROVEEDOR',
            'NRO. IMPO',
            'MONEDA',
            'NOTA CRÉDITO',
            'VALOR A PAGAR',
            'ESTADO',
            'FECHA DE VENCIMIENTO'
        ]

    def log(self, mensaje, tipo="INFO"):
        """Registra mensajes en consola y en ventana de progreso"""
        simbolos = {
            "INFO": "ℹ",
            "OK": "✓",
            "ERROR": "✗",
            "WARN": "⚠",
            "PROCESO": "►"
        }
        mensaje_formateado = f"{simbolos.get(tipo, '•')} {mensaje}"
        print(mensaje_formateado)
        
        if self.ventana_progreso:
            self.ventana_progreso.agregar_log(mensaje)
            
            if tipo == "PROCESO":
                self.ventana_progreso.actualizar_mensaje(mensaje)

    def crear_nombre_archivo(self, fecha):
        """Crea nombre del archivo basado en fecha de proyección"""
        dia = fecha.strftime('%d')
        mes = fecha.strftime('%B').upper()
        año = fecha.strftime('%Y')
        return f"{dia} {mes} {año}.xlsx"

    def crear_nombre_segunda_hoja(self, fecha):
        """Crea nombre de segunda hoja: 'MES dia'"""
        mes = fecha.strftime('%B').upper()
        dia = fecha.strftime('%d')
        return f"{mes} {dia}"
    
    def crear_estructura_carpetas(self, fecha):
        """Crea la estructura basada en fecha de proyección"""
        año_carpeta = f"AÑO {fecha.strftime('%Y')}"
        mes_carpeta = fecha.strftime('%B').upper()
        
        carpeta_destino = self.ruta_intermedio / año_carpeta / mes_carpeta
        carpeta_destino.mkdir(parents=True, exist_ok=True)
        return carpeta_destino

    def mostrar_todas_hojas(self, wb):
        """
        MÉTODO CRÍTICO: Muestra TODAS las hojas del workbook antes de copiar
        Esto evita problemas con macros que ocultan hojas (Workbook_Open, etc)
        """
        self.log("Forzando visibilidad de TODAS las hojas...", "INFO")
        try:
            for sheet in wb.Sheets:
                try:
                    # xlSheetVisible = -1
                    sheet.Visible = -1
                    self.log(f"  - Hoja '{sheet.Name}' ahora visible", "INFO")
                except Exception as e:
                    self.log(f"  - No se pudo hacer visible '{sheet.Name}': {e}", "WARN")
        except Exception as e:
            self.log(f"Error al mostrar hojas: {e}", "WARN")

    def copiar_archivo_base(self, ruta_destino):
        """
        SOLUCIÓN SIMPLE Y EFECTIVA:
        1. Abrir archivo origen
        2. Guardar como .xlsx (sin macros) → Copia TODO
        3. Eliminar hojas innecesarias
        4. Mantener solo Control_Pagos
        """
        self.log(f"Copiando archivo completo como .xlsx...", "PROCESO")
        
        excel = None
        wb = None
        
        try:
            pythoncom.CoInitialize()
            
            # Crear instancia de Excel
            excel = win32com.client.DispatchEx("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False
            excel.AutomationSecurity = 3  # Desactivar macros
            excel.EnableEvents = False     # NO ejecutar Workbook_Open
            
            # Abrir archivo origen
            self.log(f"Abriendo archivo: {self.ruta_origen.name}", "INFO")
            wb = excel.Workbooks.Open(
                str(self.ruta_origen),
                ReadOnly=True,
                UpdateLinks=0,
                IgnoreReadOnlyRecommended=True,
                Notify=False
            )
            
            # FORZAR visibilidad de todas las hojas ANTES de guardar
            self.log("Haciendo visibles todas las hojas...", "INFO")
            for sheet in wb.Sheets:
                try:
                    sheet.Visible = -1  # xlSheetVisible
                    self.log(f"  ✓ '{sheet.Name}' visible", "INFO")
                except Exception as e:
                    self.log(f"  ⚠ No se pudo hacer visible '{sheet.Name}': {e}", "WARN")
            
            # Verificar que la hoja objetivo existe
            hoja_encontrada = False
            for sheet in wb.Sheets:
                if sheet.Name.lower() == self.nombre_primera_hoja.lower():
                    hoja_encontrada = True
                    self.log(f"✓ Hoja objetivo encontrada: '{sheet.Name}'", "OK")
                    break
            
            if not hoja_encontrada:
                hojas = [s.Name for s in wb.Sheets]
                raise Exception(f"No se encontró hoja '{self.nombre_primera_hoja}'. Disponibles: {hojas}")
            
            # GUARDAR COMO .XLSX (esto copia TODO el contenido sin macros)
            ruta_dest_str = str(Path(ruta_destino).resolve())
            self.log(f"Guardando como .xlsx: {Path(ruta_destino).name}", "INFO")
            
            # FileFormat 51 = xlsx (sin macros)
            wb.SaveAs(
                Filename=ruta_dest_str,
                FileFormat=51,
                CreateBackup=False
            )
            
            self.log("✓ Archivo guardado como .xlsx", "OK")
            
            # Cerrar el archivo original
            wb.Close(SaveChanges=False)
            wb = None
            
            # ABRIR EL NUEVO ARCHIVO para limpiar hojas
            self.log("Abriendo archivo nuevo para limpieza...", "INFO")
            wb = excel.Workbooks.Open(ruta_dest_str)
            
            # ELIMINAR todas las hojas EXCEPTO la que necesitamos
            self.log("Eliminando hojas innecesarias...", "INFO")
            excel.DisplayAlerts = False  # No preguntar al eliminar
            
            hojas_a_eliminar = []
            for sheet in wb.Sheets:
                if sheet.Name.lower() != self.nombre_primera_hoja.lower():
                    hojas_a_eliminar.append(sheet.Name)
            
            for nombre_hoja in hojas_a_eliminar:
                try:
                    wb.Sheets(nombre_hoja).Delete()
                    self.log(f"  ✓ Eliminada: '{nombre_hoja}'", "INFO")
                except Exception as e:
                    self.log(f"  ⚠ No se pudo eliminar '{nombre_hoja}': {e}", "WARN")
            
            excel.DisplayAlerts = True
            
            # Verificar que quedó solo la hoja correcta
            if wb.Sheets.Count == 1:
                self.log(f"✓ Archivo limpio. Solo queda: '{wb.Sheets(1).Name}'", "OK")
            else:
                self.log(f"⚠ Advertencia: Quedaron {wb.Sheets.Count} hojas", "WARN")
            
            # GUARDAR cambios
            wb.Save()
            self.log("✓ Cambios guardados", "OK")
            
            # Cerrar
            wb.Close(SaveChanges=False)
            wb = None
            
            self.log("✓ Proceso de copia completado exitosamente", "OK")
            
        except Exception as e:
            self.log(f"ERROR al copiar archivo: {e}", "ERROR")
            import traceback
            traceback.print_exc()
            raise
            
        finally:
            # Limpieza
            if wb:
                try:
                    wb.Close(SaveChanges=False)
                except:
                    pass
            
            if excel:
                try:
                    excel.Quit()
                except:
                    pass
            
            try:
                pythoncom.CoUninitialize()
            except:
                pass

    def guardar_con_reintento(self, wb, ruta):
        """Guarda un workbook con lógica de reintento"""
        while True:
            try:
                wb.save(ruta)
                break
            except PermissionError:
                self.log(f"EL ARCHIVO ESTÁ ABIERTO: {Path(ruta).name}", "WARN")
                respuesta = messagebox.askretrycancel(
                    "Archivo Abierto",
                    f"El archivo '{Path(ruta).name}' está abierto.\n\nPor favor, ciérrelo para continuar."
                )
                if not respuesta:
                    raise Exception("Cancelado por el usuario")

    def leer_datos_control_pagos(self, ruta_archivo):
        """Lee los datos del archivo de control de pagos - VERSIÓN CORREGIDA"""
        try:
            self.log(f"Leyendo hoja '{self.nombre_primera_hoja}'...", "PROCESO")
            
            # Leer con openpyxl
            df = pd.read_excel(
                ruta_archivo, 
                sheet_name=self.nombre_primera_hoja, 
                engine='openpyxl',
                dtype=str  # CRÍTICO: Leer todo como string para evitar problemas
            )
            
            # CORRECCIÓN CRÍTICA: Limpiar nombres de columnas asegurando que sean strings
            columnas_limpias = []
            for col in df.columns:
                # Convertir a string y limpiar
                col_str = str(col).strip()
                columnas_limpias.append(col_str)
            
            df.columns = columnas_limpias
            
            self.log(f"Columnas detectadas: {df.columns.tolist()[:5]}...", "INFO")
            
            # Verificar si el DataFrame está vacío
            if df.empty:
                self.log("El archivo leído no contiene datos.", "WARN")
                return None
            
            # Renombrar columnas para estandarizar (case-insensitive)
            column_mapping = {
                '# IMPORTACION': 'NRO. IMPO',
                '#IMPORTACION': 'NRO. IMPO',
                'VALOR MONEDA ORIGEN': 'VALOR A PAGAR',
                'NOTA CREDITO': 'NOTA CRÉDITO',
                'VALOR NOTA CRÉDITO': 'NOTA CRÉDITO',
                'VALOR NOTA CREDITO': 'NOTA CRÉDITO'
            }
            
            # Aplicar renombrado (case-insensitive)
            for old_col, new_col in column_mapping.items():
                for actual_col in df.columns:
                    if actual_col.upper() == old_col.upper():
                        df.rename(columns={actual_col: new_col}, inplace=True)
                        break
            
            self.log(f"Archivo leído: {len(df)} registros totales", "OK")
            return df
            
        except Exception as e:
            self.log(f"Error al leer el archivo: {str(e)}", "ERROR")
            import traceback
            traceback.print_exc()
            return None

    def filtrar_por_fecha(self, df, fecha_filtrado):
        """Filtra registros por fecha de proyección - VERSIÓN CORREGIDA"""
        self.log(f"Filtrando por fecha de proyección: {fecha_filtrado}", "PROCESO")
        
        # CORRECCIÓN: Normalizar nombres de columnas SIN usar .str en columnas que pueden ser integers
        columnas_normalizadas = []
        for col in df.columns:
            col_normalizado = str(col).strip().upper()
            columnas_normalizadas.append(col_normalizado)
        
        df.columns = columnas_normalizadas
        
        self.log(f"Columnas disponibles: {df.columns.tolist()}", "INFO")
        
        # Buscar columna de fecha relevante
        col_fecha = None
        posibles_columnas = ['FECHA DE VENCIMIENTO', 'FECHA VENCIMIENTO', 'FECHA DE PAGO', 'FECHA PAGO']
        
        for col in posibles_columnas:
            if col in df.columns:
                col_fecha = col
                break
            
        if col_fecha:
            self.log(f"Usando columna de fecha: '{col_fecha}'", "INFO")
            
            # Convertir a datetime
            df[col_fecha] = pd.to_datetime(df[col_fecha], dayfirst=True, errors='coerce')
            
            fecha_referencia = fecha_filtrado.date() if isinstance(fecha_filtrado, datetime) else fecha_filtrado
            
            # Calcular rango de la semana (Lunes a Domingo)
            inicio_semana = fecha_referencia - timedelta(days=fecha_referencia.weekday())
            fin_semana = inicio_semana + timedelta(days=6)
            
            self.log(f"Rango de semana calculado: {inicio_semana} al {fin_semana}", "INFO")
            
            # Conteo pre-filtro
            total_registros = len(df)
            
            # Registros con fecha válida
            df_con_fecha = df.dropna(subset=[col_fecha])
            registros_con_fecha = len(df_con_fecha)
            
            # Filtrar por RANGO DE FECHA (Semana completa)
            df_fecha_match = df_con_fecha[
                (df_con_fecha[col_fecha].dt.date >= inicio_semana) & 
                (df_con_fecha[col_fecha].dt.date <= fin_semana)
            ]
            registros_fecha_match = len(df_fecha_match)
            
            self.log(f"Registros totales: {total_registros}", "INFO")
            self.log(f"Registros con fecha válida en '{col_fecha}': {registros_con_fecha}", "INFO")
            self.log(f"Registros en la semana ({inicio_semana} - {fin_semana}): {registros_fecha_match}", "INFO")
            
            if registros_fecha_match == 0:
                # Mostrar muestra de fechas para diagnóstico
                muestra_fechas = df_con_fecha[col_fecha].dt.date.unique()[:5]
                self.log(f"Muestra de fechas en el archivo: {muestra_fechas}", "WARN")
                return pd.DataFrame()

            # Filtrar por estado si existe
            if 'ESTADO' in df.columns:
                df_fecha_match = df_fecha_match.copy()
                df_fecha_match['ESTADO_NORM'] = df_fecha_match['ESTADO'].astype(str).str.upper().str.strip()
                
                # Filtrar donde contenga 'PAGAR'
                df_filtrado = df_fecha_match[
                    df_fecha_match['ESTADO_NORM'].str.contains('PAGAR', na=False)
                ].copy()
                
                registros_finales = len(df_filtrado)
                self.log(f"Registros tras filtro de estado ('PAGAR'): {registros_finales}", "INFO")
                
                if registros_finales == 0 and registros_fecha_match > 0:
                    estados_encontrados = df_fecha_match['ESTADO'].unique()
                    self.log(f"Estados encontrados: {estados_encontrados}", "WARN")
                    self.log("⚠️ No se encontraron registros con estado 'PAGAR'. Se incluirán todos los de la fecha.", "WARN")
                    df_filtrado = df_fecha_match.copy()
                
                if 'ESTADO_NORM' in df_filtrado.columns:
                    df_filtrado = df_filtrado.drop(columns=['ESTADO_NORM'])
                    
                return df_filtrado
            else:
                self.log("No se encontró columna 'ESTADO', retornando todos los registros de la fecha", "WARN")
                return df_fecha_match
        else:
            self.log(f"No se encontró columna de fecha compatible. Buscado: {posibles_columnas}", "ERROR")
            return pd.DataFrame()

    def preparar_datos_segunda_hoja(self, df_filtrado):
        """Prepara dataframe para la segunda hoja"""
        self.log(f"Preparando datos para proyección...", "PROCESO")
        
        df_resultado = pd.DataFrame()
        
        cols_map = {
            'IMPORTADOR': 'IMPORTADOR',
            'MARCA': 'MARCA',
            'PROVEEDOR': 'PROVEEDOR',
            'NRO. IMPO': 'NRO. IMPO',
            'MONEDA': 'MONEDA',
            'NOTA CRÉDITO': 'NOTA CRÉDITO',
            'VALOR A PAGAR': 'VALOR A PAGAR',
            'ESTADO': 'ESTADO'
        }
        
        for col_dest, col_origen in cols_map.items():
            if col_origen in df_filtrado.columns:
                df_resultado[col_dest] = df_filtrado[col_origen]
            else:
                df_resultado[col_dest] = ''
                
        if 'NOTA CRÉDITO' not in df_resultado.columns:
            df_resultado['NOTA CRÉDITO'] = 0.00
            
        return df_resultado

    def agrupar_y_calcular(self, df):
        """Agrupa y calcula totales"""
        self.log(f"Agrupando registros...", "PROCESO")
        
        df['VALOR A PAGAR'] = pd.to_numeric(df['VALOR A PAGAR'], errors='coerce').fillna(0)
        df = df.sort_values(by=['IMPORTADOR', 'PROVEEDOR']).reset_index(drop=True)
        
        filas_resultado = []
        grupos = df.groupby(['IMPORTADOR', 'PROVEEDOR'], sort=False)
        
        for (importador, proveedor), grupo in grupos:
            for _, registro in grupo.iterrows():
                row_dict = registro.to_dict()
                filas_resultado.append(row_dict)
            
            if len(grupo) > 1:
                total = grupo['VALOR A PAGAR'].sum()
                moneda = grupo['MONEDA'].iloc[0]
                
                fila_total = {col: '' for col in df.columns}
                fila_total['VALOR A PAGAR'] = total
                fila_total['MONEDA'] = moneda
                fila_total['_ES_TOTAL'] = True 
                filas_resultado.append(fila_total)
            
            # Agregar dos filas vacías
            fila_vacia = {col: '' for col in df.columns}
            filas_resultado.append(fila_vacia)
            filas_resultado.append(fila_vacia.copy())
        
        df_resultado = pd.DataFrame(filas_resultado)
        if '_ES_TOTAL' in df_resultado.columns:
            df_resultado = df_resultado.drop(columns=['_ES_TOTAL'])
            
        return df_resultado

    def guardar_proyeccion_com(self, ruta_archivo, df_datos, nombre_hoja):
        """Guarda la proyección usando COM"""
        self.log(f"Guardando proyección...", "PROCESO")
        
        excel = None
        wb = None
        try:
            pythoncom.CoInitialize()
            excel = win32com.client.Dispatch("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False
            excel.AutomationSecurity = 3
            
            ruta_abs = str(Path(ruta_archivo).resolve())
            wb = excel.Workbooks.Open(ruta_abs)
            
            # Crear segunda hoja
            try:
                ws = wb.Sheets.Add(After=wb.Sheets(wb.Sheets.Count))
                ws.Name = nombre_hoja
            except Exception:
                ws = excel.ActiveSheet
            
            # Preparar datos
            datos = [df_datos.columns.tolist()] + df_datos.fillna("").values.tolist()
            
            filas = len(datos)
            columnas = len(datos[0])
            
            # Escribir datos
            rango_datos = ws.Range(ws.Cells(1, 1), ws.Cells(filas, columnas))
            rango_datos.Value = datos
            
            # Formato header
            rango_header = ws.Range(ws.Cells(1, 1), ws.Cells(1, columnas))
            rango_header.Interior.Color = 11764117
            rango_header.Font.Bold = True
            rango_header.Font.Color = 16777215
            rango_header.HorizontalAlignment = -4108
            rango_header.VerticalAlignment = -4108
            rango_header.Borders.LineStyle = 1
            
            # Bordes
            rango_completo = ws.Range(ws.Cells(1, 1), ws.Cells(filas, columnas))
            rango_completo.Borders.LineStyle = 1
            
            # Formatos especiales
            col_imp = 1
            col_val = 7
            
            for i in range(2, filas + 1):
                val_imp = ws.Cells(i, col_imp).Value
                val_pago = ws.Cells(i, col_val).Value
                
                # Fila vacía
                if (val_imp is None or str(val_imp).strip() == "") and \
                   (val_pago is None or str(val_pago).strip() == ""):
                    rango_fila = ws.Range(ws.Cells(i, 1), ws.Cells(i, columnas))
                    rango_fila.Borders.LineStyle = -4142
                    rango_fila.Interior.Pattern = -4142
                
                # Fila de Total
                elif val_imp is None or str(val_imp).strip() == "":
                    rango_fila = ws.Range(ws.Cells(i, 1), ws.Cells(i, columnas))
                    rango_fila.Interior.Color = 12117678
                    rango_fila.Font.Bold = True
            
            # Formato numérico
            rango_vals = ws.Range(ws.Cells(2, col_val), ws.Cells(filas, col_val))
            rango_vals.NumberFormat = "#,##0.00"
            
            # Autofit
            ws.Columns.AutoFit()
            
            # Freeze panes
            excel.ActiveWindow.SplitRow = 1
            excel.ActiveWindow.FreezePanes = True
            
            wb.Save()
            self.log(f"Proyección guardada correctamente", "OK")
            
        except Exception as e:
            self.log(f"Error en proyección: {str(e)}", "ERROR")
            raise e
        finally:
            if wb:
                wb.Close()
            if excel:
                excel.Quit()
            pythoncom.CoUninitialize()

    def anexar_archivo_final_com(self, df_detalle):
        """Anexa registros al archivo final"""
        self.log(f"Anexando al archivo final...", "PROCESO")
        
        if not self.ruta_destino_final.exists():
            self.log("Archivo final no existe", "ERROR")
            return

        excel = None
        wb = None
        try:
            pythoncom.CoInitialize()
            excel = win32com.client.Dispatch("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False
            excel.AutomationSecurity = 3
            
            ruta_abs = str(self.ruta_destino_final.resolve())
            wb = excel.Workbooks.Open(ruta_abs)
            
            # Buscar hoja
            ws = None
            for sheet in wb.Sheets:
                sheet.Visible = -1
                if sheet.Name.lower() in ["pagos importación", "pagos importacion"]:
                    ws = sheet
                    break
            
            if not ws:
                ws = wb.ActiveSheet
            
            # Preparar datos
            datos = df_detalle.fillna("").values.tolist()
            num_nuevas_filas = len(datos)
            if num_nuevas_filas == 0:
                return

            # Encontrar última fila
            last_row = ws.Cells(ws.Rows.Count, 1).End(-4162).Row
            start_row = last_row + 1
            
            # Escribir datos
            filas = len(datos)
            columnas = len(datos[0])
            rango_dest = ws.Range(ws.Cells(start_row, 1), ws.Cells(start_row + filas - 1, columnas))
            rango_dest.Value = datos
            
            # Expandir tabla si existe
            if ws.ListObjects.Count > 0:
                tbl = ws.ListObjects(1)
                rango_tbl_header = tbl.HeaderRowRange
                fila_inicio = rango_tbl_header.Row
                col_inicio = rango_tbl_header.Column
                
                nuevo_rango_str = f"{ws.Cells(fila_inicio, col_inicio).Address}:{ws.Cells(start_row + filas - 1, columnas).Address}"
                
                try:
                    tbl.Resize(ws.Range(nuevo_rango_str))
                    self.log("Tabla expandida correctamente", "OK")
                except Exception as e:
                    self.log(f"No se pudo redimensionar tabla: {e}", "WARN")
            
            wb.Save()
            self.log("Registros anexados exitosamente", "OK")
            
        except Exception as e:
            self.log(f"Error en archivo final: {str(e)}", "ERROR")
            raise e
        finally:
            if wb:
                wb.Close()
            if excel:
                excel.Quit()
            pythoncom.CoUninitialize()
            
    def preparar_df_final(self, df_detalle):
        """Prepara DataFrame final"""
        df_final_append = pd.DataFrame()
        fecha_proyeccion = self.fecha_filtrado
        
        df_final_append['IMPORTADOR'] = df_detalle['IMPORTADOR']
        df_final_append['MARCA'] = df_detalle['MARCA']
        df_final_append['FECHA DE PAGO'] = fecha_proyeccion.strftime('%d/%m/%Y')
        df_final_append['DIA'] = fecha_proyeccion.day
        df_final_append['MES'] = fecha_proyeccion.month
        df_final_append['AÑO'] = fecha_proyeccion.year
        df_final_append['PROVEEDOR'] = df_detalle['PROVEEDOR']
        df_final_append['# IMPORTACION'] = df_detalle['NRO. IMPO']
        df_final_append['VALOR MONEDA ORIGEN'] = df_detalle['VALOR A PAGAR']
        df_final_append['MONEDA'] = df_detalle['MONEDA']
        
        def calc_valor_usd(row):
            if str(row['MONEDA']).upper() == 'USD':
                return row['VALOR A PAGAR']
            return ''
            
        def calc_factor(row):
            if str(row['MONEDA']).upper() == 'USD':
                return 1
            return ''
            
        df_final_append['VALOR USD'] = df_detalle.apply(calc_valor_usd, axis=1)
        df_final_append['FACTOR DE CONVERSION'] = df_detalle.apply(calc_factor, axis=1)
        df_final_append['DESCUENTO PRONTO PAGO'] = 0
        df_final_append['FORMA DE PAGO'] = ''
        df_final_append['TIPO DE PAGO'] = 'CUENTA COMPENSACION'
        df_final_append['FECHA DE APERTURA CREDITO -UTILIZACION LC'] = 'N/A'
        df_final_append['FECHA DE VENCIMIENTO'] = 'N/A'
        df_final_append['# CREDITO'] = 'N/A'
        df_final_append['# DEUDA EXTERNA'] = 'N/A'
        df_final_append['NOTA CREDITO'] = 0.00
        df_final_append['OBSERVACIONES'] = ''
        
        return df_final_append

    def agregar_a_archivo_final(self, df_detalle):
        """Agrega registros al archivo final"""
        try:
            df_final = self.preparar_df_final(df_detalle)
            self.anexar_archivo_final_com(df_final)
        except Exception as e:
            self.log(f"Error en proceso final: {str(e)}", "ERROR")

    def ejecutar_proceso(self):
        """Ejecuta el proceso completo"""
        print("\n" + "="*80)
        print("    AUTOMATIZACIÓN DE CONTROL DE PAGOS - VERSIÓN 2.0 CORREGIDA")
        print("="*80 + "\n")
        
        try:
            if not self.ruta_origen.exists():
                self.log(f"No se encuentra el archivo original", "ERROR")
                messagebox.showerror("Error", f"No se encuentra el archivo:\n{self.ruta_origen}")
                return None
            
            fecha_proyeccion = self.fecha_filtrado
            self.log(f"Fecha de proyección: {fecha_proyeccion.strftime('%d/%m/%Y')}", "INFO")
            
            carpeta_destino = self.crear_estructura_carpetas(fecha_proyeccion)
            nombre_archivo = self.crear_nombre_archivo(fecha_proyeccion)
            ruta_archivo_nuevo = carpeta_destino / nombre_archivo
            
            self.copiar_archivo_base(ruta_archivo_nuevo)
            
            df_original = self.leer_datos_control_pagos(ruta_archivo_nuevo)
            if df_original is None:
                return None
            
            df_filtrado = self.filtrar_por_fecha(df_original, fecha_proyeccion)
            
            if len(df_filtrado) == 0:
                self.log("No se encontraron registros", "WARN")
                messagebox.showwarning("Sin registros", "No se encontraron registros para la fecha seleccionada.")
                return
            
            df_segunda = self.preparar_datos_segunda_hoja(df_filtrado)
            df_agrupado = self.agrupar_y_calcular(df_segunda)
            
            nombre_segunda_hoja = self.crear_nombre_segunda_hoja(fecha_proyeccion)
            
            self.guardar_proyeccion_com(ruta_archivo_nuevo, df_agrupado, nombre_segunda_hoja)
            
            self.agregar_a_archivo_final(df_segunda)
            
            print("\n" + "="*80)
            print("PROCESO COMPLETADO EXITOSAMENTE")
            print("="*80)
            
            messagebox.showinfo(
                "¡Proceso Completado!",
                f"El proceso ha finalizado exitosamente.\n\n"
                f"📁 Proyección guardada en:\n{ruta_archivo_nuevo}\n\n"
                f"📁 Archivo final actualizado:\n{self.ruta_destino_final.name}"
            )
            return str(ruta_archivo_nuevo)
            
        except Exception as e:
            self.log(f"ERROR CRÍTICO: {str(e)}", "ERROR")
            import traceback
            traceback.print_exc()
            messagebox.showerror("Error", f"Ocurrió un error:\n\n{str(e)}")
            return None

def main():
    """Función principal de la aplicación"""
    # Mostrar ventana de selección de fecha
    interfaz = InterfazModerna()
    interfaz.crear_ventana()
    
    if not interfaz.ejecutar_proceso:
        return
    
    # Mensaje de confirmación
    if not messagebox.askyesno(
        "Confirmar Ejecución",
        "Antes de continuar, asegúrese de:\n\n"
        "✓ Haber actualizado el archivo 'CONTROL DE PAGOS.xlsm'\n"
        "✓ Haber guardado todos los cambios\n"
        "✓ Cerrar el archivo si está abierto\n\n"
        "¿Desea continuar?"
    ):
        return
    
    # Crear ventana de progreso
    ventana_prog = VentanaProgreso()
    
    try:
        # Ejecutar proceso
        copiador = CopiarArchivo(
            fecha_filtrado=interfaz.fecha_seleccionada,
            ventana_progreso=ventana_prog
        )
        resultado = copiador.ejecutar_proceso()
        
        # Cerrar ventana de progreso
        ventana_prog.cerrar()
        
    except Exception as e:
        ventana_prog.cerrar()
        messagebox.showerror("Error Fatal", f"Error inesperado:\n\n{str(e)}")

if __name__ == "__main__":
    main()